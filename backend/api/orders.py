"""Requester order submission, history, self-editing and count requests."""
import asyncio
from datetime import datetime, timezone
import csv, io, json, html

from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File
from fastapi.responses import Response, HTMLResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_

from backend.db.session import get_db
from backend.schemas.schemas import (
    OrderCreate, OrderOut, OrdersUpdateResponse, PendingOrderUpdate,
    CountRequestCreate, CountRequestOut,
)
from backend.models.models import (
    Order, OrderLine, User, OrderProofPhoto, CountRequest, ProjectMember, OrderCollaborator, OrderAttendeeList, Kit,
)
from backend.auth.dependencies import get_current_user, get_current_user_flexible
from backend.services import order_service, item_service, access_service

router = APIRouter(tags=["orders"])
POLL_INTERVAL_SECONDS = 1.0
POLL_MAX_WAIT_SECONDS = 25.0


def _my_orders_query(db: Session, user_id: int):
    collaborator_orders = db.query(OrderCollaborator.order_id).filter(
        OrderCollaborator.user_id == user_id)
    member_projects = db.query(ProjectMember.project_id).filter(ProjectMember.user_id == user_id)
    return (db.query(Order).options(
        joinedload(Order.lines).joinedload(OrderLine.item),
        joinedload(Order.project), joinedload(Order.requester),
        joinedload(Order.tracking_numbers), joinedload(Order.proof_photos),
        joinedload(Order.collaborators).joinedload(OrderCollaborator.user),
        joinedload(Order.kit_selections), joinedload(Order.attendee_lists),
    ).filter(or_(Order.requester_user_id == user_id,
                 Order.id.in_(collaborator_orders),
                 Order.project_id.in_(member_projects))))


def _count_out(row: CountRequest) -> CountRequestOut:
    return CountRequestOut(
        id=row.id, item_id=row.item_id, item_code=row.item.code,
        item_name=row.item.name,
        requester=row.requester.full_name or row.requester.username,
        note=row.note or "", status=row.status,
        resolution_note=row.resolution_note or "",
        created_at=row.created_at, resolved_at=row.resolved_at,
        system_qty_before=row.system_qty_before, physical_qty=row.physical_qty,
        adjustment_delta=row.adjustment_delta,
    )


@router.post("/orders", response_model=OrderOut, status_code=201)
def create_order(body: OrderCreate, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    order = order_service.create_order(
        db, requester=user, project_id=body.project_id,
        project_data=(body.new_project.model_dump() if body.new_project else None),
        save_project=body.save_project, notes=body.notes, lines=body.lines, kits=body.kits,
        collaborator_user_ids=body.collaborator_user_ids, source="browser")
    return order_service.to_order_out(order, viewer=user, db=db)


@router.get("/orders/my", response_model=list[OrderOut])
def my_orders(db: Session = Depends(get_db),
              user: User = Depends(get_current_user)):
    orders = (_my_orders_query(db, user.id)
              .filter(Order.deleted_at.is_(None))
              .order_by(Order.created_at.desc()).all())
    return [order_service.to_order_out(o, viewer=user, db=db) for o in orders]


@router.put("/orders/{order_id}", response_model=OrderOut)
def edit_my_pending_order(order_id: int, body: PendingOrderUpdate,
                          db: Session = Depends(get_db),
                          user: User = Depends(get_current_user)):
    order = order_service.edit_order(
        db, order_id=order_id, actor=user, notes=body.notes,
        project_data=(body.project.model_dump(exclude_none=True)
                      if body.project else None),
        lines=body.lines, collaborator_user_ids=body.collaborator_user_ids,
        owner_only_pending=True, source="browser")
    return order_service.to_order_out(order, viewer=user, db=db)


@router.post("/orders/{order_id}/cancel", response_model=OrderOut)
def cancel_my_order(order_id: int, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)):
    order = order_service.cancel_order(db, order_id=order_id, actor=user, source="browser")
    return order_service.to_order_out(order, viewer=user, db=db)


@router.post("/catalog/{item_id}/count-requests",
             response_model=CountRequestOut, status_code=201)
def request_count(item_id: int, body: CountRequestCreate,
                  db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    row = item_service.create_count_request(
        db, item_id=item_id, requester=user, note=body.note, source="browser")
    return _count_out(row)


@router.get("/order-proof/{photo_id}", include_in_schema=False)
def order_proof_photo(photo_id: int, db: Session = Depends(get_db),
                      user: User = Depends(get_current_user)):
    photo = (db.query(OrderProofPhoto).options(joinedload(OrderProofPhoto.order))
             .filter_by(id=photo_id).first())
    if not photo:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Proof photo not found.")
    is_owner = photo.order.requester_user_id == user.id
    is_reviewer = user.has_role("admin") or user.has_role("approver")
    is_collaborator = db.query(OrderCollaborator).filter_by(order_id=photo.order.id, user_id=user.id).first() is not None
    is_member = bool(photo.order.project_id and access_service.project_membership(db, photo.order.project_id, user.id))
    if not (is_owner or is_reviewer or is_collaborator or is_member):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Not your order proof.")
    return Response(content=photo.content, media_type=photo.content_type,
                    headers={"Cache-Control": "private, max-age=3600"})



def _read_attendee_file(filename: str, content: bytes):
    name=(filename or "attendees.csv").lower()
    rows=[]
    if name.endswith(".csv"):
        text_content=content.decode("utf-8-sig", errors="replace")
        reader=csv.DictReader(io.StringIO(text_content))
        headers=[(h or "").strip() for h in (reader.fieldnames or [])]
        rows=[{str(k or "").strip(): (v or "") for k,v in r.items()} for r in reader]
    elif name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(500,"Excel support is not installed on the server.") from exc
        wb=load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws=wb.active
        values=list(ws.iter_rows(values_only=True))
        headers=[str(v or "").strip() for v in (values[0] if values else [])]
        for row in values[1:]:
            if any(v not in (None,"") for v in row):
                rows.append({headers[i] if i<len(headers) else f"Column {i+1}": ("" if v is None else str(v)) for i,v in enumerate(row)})
    else:
        raise HTTPException(400,"Upload a CSV or XLSX attendee list.")
    headers=[h for h in headers if h]
    aliases={"first name":"First Name","firstname":"First Name","first":"First Name","last name":"Last Name","lastname":"Last Name","last":"Last Name","company":"Company / Organization","organization":"Company / Organization","company/organization":"Company / Organization","title":"Title","job title":"Title"}
    recognized=[]
    for h in headers:
        canonical=aliases.get(h.lower().strip())
        if canonical and canonical not in recognized: recognized.append(canonical)
    return headers, rows, recognized

@router.post("/attendee-lists/preview")
async def preview_attendee_list(file: UploadFile=File(...), user: User=Depends(get_current_user)):
    content=await file.read()
    if not content or len(content)>8*1024*1024:
        raise HTTPException(400,"Attendee list must be between 1 byte and 8 MB.")
    headers, rows, recognized=_read_attendee_file(file.filename or "attendees.csv", content)
    return {"filename":file.filename or "attendees.csv","row_count":len(rows),"columns":headers,"recognized_fields":recognized,"missing_recommended":[x for x in ["First Name","Last Name"] if x not in recognized],"sample":rows[:5]}

@router.post("/orders/{order_id}/attendee-list", status_code=201)
async def upload_attendee_list(order_id:int, file:UploadFile=File(...), db:Session=Depends(get_db), user:User=Depends(get_current_user)):
    order=db.query(Order).filter(Order.id==order_id, Order.deleted_at.is_(None)).first()
    if not order: raise HTTPException(404,"Order not found.")
    if order.requester_user_id!=user.id and not user.has_role("admin"):
        raise HTTPException(403,"You cannot attach a list to this order.")
    content=await file.read()
    if not content or len(content)>8*1024*1024: raise HTTPException(400,"Attendee list must be between 1 byte and 8 MB.")
    headers, rows, recognized=_read_attendee_file(file.filename or "attendees.csv", content)
    row=OrderAttendeeList(order_id=order.id, filename=(file.filename or "attendees.csv")[:255], content_type=file.content_type or "application/octet-stream", content=content, row_count=len(rows), columns_json=json.dumps(headers))
    db.add(row); db.commit(); db.refresh(row)
    return {"id":row.id,"filename":row.filename,"row_count":row.row_count,"columns":headers,"recognized_fields":recognized}

@router.get("/orders/{order_id}/attendee-list/{attachment_id}")
def download_attendee_list(order_id:int, attachment_id:int, db:Session=Depends(get_db), user:User=Depends(get_current_user_flexible)):
    order=db.query(Order).filter_by(id=order_id).first()
    row=db.query(OrderAttendeeList).filter_by(id=attachment_id,order_id=order_id).first()
    if not order or not row: raise HTTPException(404,"Attendee list not found.")
    allowed=order.requester_user_id==user.id or user.has_role("admin") or user.has_role("approver") or db.query(OrderCollaborator).filter_by(order_id=order_id,user_id=user.id).first() is not None
    if not allowed: raise HTTPException(403,"Not allowed.")
    safe=(row.filename or "attendees.csv").replace('"','')
    return Response(content=row.content, media_type=row.content_type or "application/octet-stream", headers={"Content-Disposition":f'attachment; filename="{safe}"'})

@router.get("/orders/{order_id}/picking-sheet", response_class=HTMLResponse, include_in_schema=False)
def picking_sheet(order_id:int, db:Session=Depends(get_db), user:User=Depends(get_current_user_flexible)):
    if not (user.has_role("admin") or user.has_role("approver")):
        raise HTTPException(403,"Warehouse access required.")
    order=order_service._load_order(db,order_id)
    folder_qty={}
    folder_sections=[]
    for selection in order.kit_selections:
        if (selection.kit_type_snapshot or "standard")!="folder": continue
        kit=selection.kit
        design={}
        try: design=json.loads(selection.design_json_snapshot or "{}")
        except Exception: pass
        components=[]
        if kit:
            for c in sorted(kit.components,key=lambda x:x.position):
                total=c.quantity*selection.quantity
                folder_qty[c.item_id]=folder_qty.get(c.item_id,0)+total
                meta=next((x for x in design.get("items",[]) if str(x.get("item_id",""))==str(c.item_id) or str(x.get("sku",""))==str(c.item.code)),{})
                components.append((c.item.code,c.item.name,total,c.position,meta))
        folder_sections.append((selection,design,components))
    main_rows=[]
    for line in order.lines:
        qty=(line.qty_approved if line.qty_approved is not None else line.qty_requested)-folder_qty.get(line.item_id,0)
        if qty>0: main_rows.append((line.item_code_snapshot or line.item.code,line.item_name_snapshot or line.item.name,qty,line.item_location_snapshot or line.item.location))
    project=order.project
    attn=html.escape(getattr(project,"shipping_attention","") or "") if project else ""
    rows=''.join(f'<tr><td>{html.escape(str(c))}</td><td>{html.escape(str(n))}</td><td>{q}</td><td>{html.escape(str(loc or ""))}</td></tr>' for c,n,q,loc in main_rows) or '<tr><td colspan="4">Folder kit only. See build sheet below.</td></tr>'
    sheets=[]
    for sel,design,components in folder_sections:
        comp_rows=[]
        for i,(code,name,qty,pos,meta) in enumerate(components):
            try: x=float(meta.get("x",1))
            except Exception: x=1
            placement=meta.get("placement") or ("Back of cover" if x<.5 else "Inside right")
            layer=meta.get("layer") or meta.get("z") or pos+1
            comp_rows.append(f'<tr><td>{i+1}</td><td>{html.escape(code)}</td><td>{html.escape(name)}</td><td>{qty}</td><td>{html.escape(str(placement))}</td><td>{html.escape(str(layer))}</td></tr>')
        sheets.append(f'<section class="page folder"><h1>Folder Kit Build Sheet</h1><h2>{html.escape(sel.kit_code_snapshot)} · {html.escape(sel.kit_name_snapshot)} × {sel.quantity}</h2><p>Build components in the sequence and placement shown below. Do not mix these part numbers into the main picking list.</p><table><thead><tr><th>Step</th><th>Part #</th><th>Item</th><th>Qty</th><th>Placement</th><th>Layer</th></tr></thead><tbody>{"".join(comp_rows)}</tbody></table></section>')
    page = '<!doctype html><html><head><meta charset="utf-8"><title>'+html.escape(order.order_number or str(order.id))+' Picking</title><style>@page{size:letter;margin:.45in}body{font-family:Arial;color:#111;margin:0}.page{page-break-after:always}.page:last-child{page-break-after:auto}h1{font-size:20px}h2{font-size:14px}table{width:100%;border-collapse:collapse;margin-top:14px}th,td{border:1px solid #777;padding:7px;text-align:left;font-size:11px}th{background:#eee}.meta{display:grid;grid-template-columns:1fr 1fr;gap:5px;font-size:11px}.folder{border-top:8px solid #f2b705;padding-top:10px}</style></head><body>'
    page += f'<section class="page"><h1>Picking Slip · {html.escape(order.order_number or str(order.id))}</h1><div class="meta"><div><b>Project:</b> {html.escape(order.project.name if order.project else "")}</div><div><b>Requester:</b> {html.escape(order.requester.full_name or order.requester.username)}</div><div><b>ATTN:</b> {attn}</div><div><b>Folder sheets:</b> {len(folder_sections)}</div></div><table><thead><tr><th>Part #</th><th>Item</th><th>Qty</th><th>Loc.</th></tr></thead><tbody>{rows}</tbody></table></section>'
    page += ''.join(sheets) + '</body></html>'
    return page

@router.get("/orders/updates", response_model=OrdersUpdateResponse)
async def poll_my_orders(
    since: datetime | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if since is not None and since.tzinfo is None:
        since = since.replace(tzinfo=timezone.utc)

    async def _check():
        q = _my_orders_query(db, user.id)
        if since is not None:
            q = q.filter(Order.updated_at > since)
        return q.order_by(Order.updated_at.desc()).all()

    if since is None:
        orders = await _check()
        return OrdersUpdateResponse(
            server_time=datetime.now(timezone.utc),
            orders=[order_service.to_order_out(o, viewer=user, db=db) for o in orders])
    waited = 0.0
    while waited < POLL_MAX_WAIT_SECONDS:
        orders = await _check()
        if orders:
            return OrdersUpdateResponse(
                server_time=datetime.now(timezone.utc),
                orders=[order_service.to_order_out(o, viewer=user, db=db) for o in orders])
        await asyncio.sleep(POLL_INTERVAL_SECONDS)
        waited += POLL_INTERVAL_SECONDS
    return OrdersUpdateResponse(server_time=datetime.now(timezone.utc), orders=[])


@router.get("/orders/{order_id}", response_model=OrderOut)
def get_order(order_id: int, db: Session = Depends(get_db),
              user: User = Depends(get_current_user)):
    order = (db.query(Order).options(
        joinedload(Order.lines).joinedload(OrderLine.item),
        joinedload(Order.project), joinedload(Order.requester),
        joinedload(Order.tracking_numbers), joinedload(Order.proof_photos),
    ).filter(Order.id == order_id, Order.deleted_at.is_(None)).first())
    if not order:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Order not found.")
    is_owner = order.requester_user_id == user.id
    is_reviewer = user.has_role("admin") or user.has_role("approver")
    is_member = bool(order.project_id and
                     access_service.project_membership(db, order.project_id, user.id))
    if not (is_owner or is_reviewer or is_member):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Not your order.")
    return order_service.to_order_out(order, viewer=user, db=db)
