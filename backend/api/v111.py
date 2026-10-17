"""v1.11 endpoints: item submissions + review, folder-kit management, saved
addresses, order printouts, Badge Studio storage, Business Development brand
fix, and the admin reports (order timing + demand / dead stock).

Everything here is additive: older admin apps and pages keep working because
no existing route changed its shape.
"""
from __future__ import annotations

import io
import json
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload

from backend.auth.dependencies import get_current_user, get_current_user_flexible, require_role
from backend.db.session import get_db
from backend.models.models import (
    Approval, BadgeProject, InventoryTransaction, Item, Kit, KitComponent, Order,
    OrderCollaborator, OrderLine, OrderPrintFile, Project, User,
)
from backend.schemas.schemas import (
    BrandFixRequest, FolderKitUpdate, ItemOut, ItemReviewDecision, ItemSubmission,
    PrintFileUpdate,
)
from backend.services import access_service, item_service
from backend.services.audit_service import log_action
from backend.services.v111_service import (
    item_out_kwargs, notify_admins, notify_user, refresh_kit_review,
    refresh_kits_for_item, suggest_code, utcnow,
)

router = APIRouter(tags=["v1.11"])

MAX_PRINT_BYTES = 25 * 1024 * 1024
PRINT_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MAX_BADGE_DOC_CHARS = 3_000_000


def _aware(dt):
    if dt is None:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


# ---------------------------------------------------------------------------
# Item submissions (Folder Builder "New item") and admin review
# ---------------------------------------------------------------------------
@router.get("/catalog/suggest-code")
def suggest_part_number(prefix: str = Query("LIT", max_length=20), db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    return {"code": suggest_code(db, prefix)}


@router.post("/catalog/item-submissions", response_model=ItemOut, status_code=201)
def submit_item(body: ItemSubmission, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    code = (body.code or "").strip().upper() or suggest_code(db, "LIT")
    if db.query(Item.id).filter(func.lower(Item.code) == code.lower()).first():
        raise HTTPException(status.HTTP_409_CONFLICT, f"Part number {code} already exists. Pick another.")
    is_admin = user.has_role("admin")
    data = {
        "code": code, "name": body.name.strip(), "description": body.description.strip(),
        "category": body.category.strip(), "brand": body.brand.strip(), "measures": body.measures.strip(),
        "location": "", "qty_on_hand": None, "nav_tracked": False, "nav_item_number": "",
        "reorder_threshold": 0, "cost": 0.0, "active": True,
    }
    item = item_service.create_item(db, data=data, actor=user, source="browser")
    item.created_by_user_id = user.id
    # Admins adding from the Folder Builder still go through review so stock
    # and photos get checked, but they can approve it themselves right away.
    item.review_status = "pending_review"
    notify_admins(db, "item_review", f"New item to review: {item.code}",
                  f"{user.full_name or user.username} added \"{item.name}\" from the Folder Builder."
                  + (" (added by an admin)" if is_admin else ""), "item", item.id)
    log_action(db, user_id=user.id, action="item.submit", object_type="item", object_id=item.id,
               new_value={"code": item.code, "name": item.name}, source="browser")
    db.commit(); db.refresh(item)
    return ItemOut.from_orm_item(item)


def _own_submission(db: Session, item_id: int, user: User) -> Item:
    item = db.query(Item).filter(Item.id == item_id, Item.deleted_at.is_(None)).first()
    if not item or (item.created_by_user_id != user.id and not user.has_role("admin")):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found.")
    return item


@router.post("/catalog/item-submissions/{item_id}/image", response_model=ItemOut, status_code=201)
async def submit_item_image(item_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                            user: User = Depends(get_current_user)):
    item = _own_submission(db, item_id, user)
    if (item.review_status or "live") == "live" and not user.has_role("admin"):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This item is already live. Ask an admin to change its photo.")
    item_service.add_item_image(db, item_id=item.id, filename=file.filename or "photo.png",
                                content=await file.read(), actor=user, source="browser")
    db.refresh(item)
    return ItemOut.from_orm_item(item)


@router.get("/catalog/item-submissions/mine", response_model=list[ItemOut])
def my_submissions(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = (db.query(Item).filter(Item.created_by_user_id == user.id, Item.deleted_at.is_(None),
                                  Item.review_status.in_(["pending_review", "rejected"]))
            .order_by(Item.created_at.desc()).all())
    return [ItemOut.from_orm_item(i) for i in rows]


@router.get("/admin/item-reviews")
def item_reviews(db: Session = Depends(get_db), _u: User = Depends(require_role("admin"))):
    rows = (db.query(Item).filter(Item.deleted_at.is_(None), Item.review_status == "pending_review")
            .order_by(Item.created_at).all())
    users = {u.id: (u.full_name or u.username) for u in db.query(User).all()}
    out = []
    for i in rows:
        kits = (db.query(Kit).join(KitComponent, KitComponent.kit_id == Kit.id)
                .filter(KitComponent.item_id == i.id).all())
        d = ItemOut.from_orm_item(i).model_dump()
        d["submitted_by"] = users.get(i.created_by_user_id, "")
        d["submitted_at"] = i.created_at.isoformat() if i.created_at else ""
        d["used_in_kits"] = [{"id": k.id, "code": k.code, "name": k.name} for k in kits]
        out.append(d)
    return out


@router.post("/admin/items/{item_id}/approve", response_model=ItemOut)
def approve_item(item_id: int, body: ItemReviewDecision, db: Session = Depends(get_db),
                 user: User = Depends(require_role("admin"))):
    item = db.query(Item).filter(Item.id == item_id, Item.deleted_at.is_(None)).first()
    if not item:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found.")
    item.review_status = "live"
    item.review_note = (body.note or "").strip()
    item.active = True
    db.flush()
    if body.opening_qty is not None:
        from backend.models.models import CountRequest
        open_count = (db.query(CountRequest).filter_by(item_id=item.id, status="open")
                      .order_by(CountRequest.id).first())
        if open_count is None:
            open_count = CountRequest(item_id=item.id, requester_user_id=user.id,
                                      note="Initial count (item review)", status="open")
            db.add(open_count); db.flush()
        db.commit()
        item_service.resolve_count_request(db, request_id=open_count.id, actor=user,
                                           physical_quantity=int(body.opening_qty),
                                           resolution_note="Counted during item review",
                                           inventory_location="0", source="admin_app")
        item = db.query(Item).filter_by(id=item_id).first()
    changed = refresh_kits_for_item(db, item.id)
    notify_user(db, item.created_by_user_id, "item_approved", f"{item.code} is live",
                f"\"{item.name}\" was approved and is now in the catalog."
                + (f" Note: {item.review_note}" if item.review_note else ""), "item", item.id)
    for kit in changed:
        if kit.review_status == "live":
            notify_user(db, kit.created_by_user_id, "kit_live", f"Folder kit {kit.code} is live",
                        "Every item in this folder kit has been approved.", "kit", kit.id)
    log_action(db, user_id=user.id, action="item.approve", object_type="item", object_id=item.id,
               new_value={"note": item.review_note, "kits_now_live": [k.code for k in changed if k.review_status == "live"]},
               source="admin_app")
    db.commit(); db.refresh(item)
    return ItemOut.from_orm_item(item)


@router.post("/admin/items/{item_id}/reject", response_model=ItemOut)
def reject_item(item_id: int, body: ItemReviewDecision, db: Session = Depends(get_db),
                user: User = Depends(require_role("admin"))):
    item = db.query(Item).filter(Item.id == item_id, Item.deleted_at.is_(None)).first()
    if not item:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found.")
    item.review_status = "rejected"
    item.review_note = (body.note or "").strip()
    item.active = False
    db.flush()
    refresh_kits_for_item(db, item.id)
    notify_user(db, item.created_by_user_id, "item_rejected", f"{item.code} was not added",
                item.review_note or "An admin declined this item.", "item", item.id)
    log_action(db, user_id=user.id, action="item.reject", object_type="item", object_id=item.id,
               new_value={"note": item.review_note}, source="admin_app")
    db.commit(); db.refresh(item)
    return ItemOut.from_orm_item(item)


# ---------------------------------------------------------------------------
# Business Development (BDG) brand fix — older builds mapped BDG to Bernafon.
# Admin picks exactly which items to move; nothing moves automatically.
# ---------------------------------------------------------------------------
@router.get("/admin/brand-fix/candidates", response_model=list[ItemOut])
def brand_fix_candidates(brand: str = Query("Bernafon"), db: Session = Depends(get_db),
                         _u: User = Depends(require_role("admin"))):
    rows = (db.query(Item).filter(Item.deleted_at.is_(None),
                                  func.lower(func.trim(Item.brand)) == brand.strip().lower())
            .order_by(Item.code).all())
    return [ItemOut.from_orm_item(i) for i in rows]


@router.post("/admin/brand-fix")
def brand_fix(body: BrandFixRequest, db: Session = Depends(get_db),
              user: User = Depends(require_role("admin"))):
    rows = db.query(Item).filter(Item.id.in_(body.item_ids), Item.deleted_at.is_(None)).all()
    moved = []
    for item in rows:
        old = item.brand
        item.brand = body.to_brand.strip()
        moved.append({"id": item.id, "code": item.code, "from": old, "to": item.brand})
    log_action(db, user_id=user.id, action="item.brand_fix", object_type="item", object_id="bulk",
               new_value=moved, source="admin_app")
    db.commit()
    return {"moved": len(moved), "items": moved}


# ---------------------------------------------------------------------------
# Folder kits: list mine, edit, duplicate, retire
# ---------------------------------------------------------------------------
def _kit_dict(db: Session, kit: Kit) -> dict:
    warnings = []
    comps = []
    for c in sorted(kit.components, key=lambda x: x.position):
        item = c.item
        if item is None or item.deleted_at is not None:
            warnings.append(f"Item {c.item_id} was deleted")
        elif (item.review_status or "live") == "pending_review":
            warnings.append(f"{item.code} is waiting for admin review")
        elif (item.review_status or "live") == "rejected" or not item.active:
            warnings.append(f"{item.code} is no longer available")
        elif int(item.qty_on_hand or 0) < int(c.quantity or 1):
            warnings.append(f"{item.code} is out of stock" if (item.qty_on_hand or 0) <= 0
                            else f"{item.code} is low ({item.qty_on_hand} left)")
        comps.append({"item_id": c.item_id, "quantity": c.quantity, "position": c.position,
                      "item_code": item.code if item else "", "item_name": item.name if item else "Deleted item",
                      "image_id": item.images[0].id if item and item.images else None,
                      "review_status": (item.review_status or "live") if item else "deleted"})
    return {"id": kit.id, "name": kit.name, "code": kit.code, "description": kit.description or "",
            "design_json": kit.design_json or "", "active": kit.active,
            "review_status": kit.review_status or "live", "image_available": bool(kit.image),
            "created_by_user_id": kit.created_by_user_id,
            "updated_at": kit.updated_at.isoformat() if kit.updated_at else "",
            "warnings": warnings, "components": comps}


def _editable_kit(db: Session, kit_id: int, user: User) -> Kit:
    kit = db.query(Kit).filter_by(id=kit_id, kit_type="folder").first()
    if not kit or (kit.created_by_user_id != user.id and not user.has_role("admin")):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Folder kit not found.")
    return kit


@router.get("/folder-kits/mine")
def my_folder_kits(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(Kit).filter(Kit.kit_type == "folder", Kit.active.is_(True))
    if not user.has_role("admin"):
        q = q.filter(Kit.created_by_user_id == user.id)
    return [_kit_dict(db, k) for k in q.order_by(Kit.updated_at.desc()).all()]


@router.put("/folder-kits/{kit_id}")
def update_folder_kit(kit_id: int, body: FolderKitUpdate, db: Session = Depends(get_db),
                      user: User = Depends(get_current_user)):
    from backend.api.catalog import can_use_in_folder
    kit = _editable_kit(db, kit_id, user)
    if body.code is not None:
        code = body.code.strip()
        clash = db.query(Kit).filter(func.lower(Kit.code) == code.lower(), Kit.id != kit.id).first()
        if clash:
            raise HTTPException(status.HTTP_409_CONFLICT, "That kit code already exists.")
        kit.code = code
    if body.name is not None:
        kit.name = body.name.strip()
    if body.description is not None:
        kit.description = body.description
    if body.design_json is not None:
        kit.design_json = body.design_json
    if body.active is not None:
        kit.active = body.active
    if body.components is not None:
        if not body.components:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Add at least one catalog item to the folder.")
        kit.components.clear()
        db.flush()
        for i, c in enumerate(body.components):
            item = db.query(Item).filter_by(id=c.item_id).first()
            if not item or not can_use_in_folder(user, item):
                raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Item {c.item_id} is not available.")
            kit.components.append(KitComponent(item_id=item.id, quantity=c.quantity,
                                               position=c.position if c.position is not None else i))
    db.flush()
    refresh_kit_review(db, kit)
    log_action(db, user_id=user.id, action="kit.update", object_type="kit", object_id=kit.id,
               new_value=body.model_dump(exclude_none=True, exclude={"design_json"}), source="browser")
    db.commit(); db.refresh(kit)
    return _kit_dict(db, kit)


@router.post("/folder-kits/{kit_id}/duplicate", status_code=201)
def duplicate_folder_kit(kit_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    src = db.query(Kit).filter_by(id=kit_id, kit_type="folder").first()
    if not src:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Folder kit not found.")
    if src.created_by_user_id != user.id and not user.has_role("admin") and (src.review_status or "live") != "live":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Folder kit not found.")
    base = f"{src.code}-COPY"
    code, n = base, 2
    while db.query(Kit.id).filter(func.lower(Kit.code) == code.lower()).first():
        code = f"{base}{n}"; n += 1
    kit = Kit(name=f"{src.name} (copy)"[:200], code=code[:60], brand=src.brand or "",
              description=src.description or "", active=True, custom=True, saved_for_reuse=True,
              kit_type="folder", design_json=src.design_json or "", created_by_user_id=user.id)
    db.add(kit); db.flush()
    for c in src.components:
        kit.components.append(KitComponent(item_id=c.item_id, quantity=c.quantity, position=c.position))
    if src.image:
        from backend.models.models import KitImage
        db.add(KitImage(kit_id=kit.id, filename=src.image.filename, content_type=src.image.content_type,
                        content=src.image.content))
    db.flush()
    refresh_kit_review(db, kit)
    log_action(db, user_id=user.id, action="kit.duplicate", object_type="kit", object_id=kit.id,
               new_value={"from": src.code, "code": kit.code}, source="browser")
    db.commit(); db.refresh(kit)
    return _kit_dict(db, kit)


@router.delete("/folder-kits/{kit_id}", status_code=204)
def retire_folder_kit(kit_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    kit = _editable_kit(db, kit_id, user)
    kit.active = False
    log_action(db, user_id=user.id, action="kit.retire", object_type="kit", object_id=kit.id,
               new_value={"code": kit.code}, source="browser")
    db.commit()
    return None


# ---------------------------------------------------------------------------
# Saved addresses — distinct delivery addresses from the user's own orders
# ---------------------------------------------------------------------------
@router.get("/me/addresses")
def my_addresses(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = (db.query(Project, Order.created_at).join(Order, Order.project_id == Project.id)
            .filter(Order.requester_user_id == user.id, Order.deleted_at.is_(None))
            .order_by(Order.created_at.desc()).limit(300).all())
    seen, out = set(), []
    for p, created in rows:
        if not (p.shipping_address1 and p.shipping_city and p.shipping_state):
            continue
        key = tuple((x or "").strip().lower() for x in (p.shipping_address1, p.shipping_address2,
                                                         p.shipping_city, p.shipping_state, p.shipping_postal_code))
        if key in seen:
            continue
        seen.add(key)
        out.append({"label": f"{p.shipping_address1}, {p.shipping_city}, {p.shipping_state}",
                    "shipping_address1": p.shipping_address1 or "", "shipping_address2": p.shipping_address2 or "",
                    "shipping_attention": p.shipping_attention or "", "shipping_city": p.shipping_city or "",
                    "shipping_state": p.shipping_state or "", "shipping_postal_code": p.shipping_postal_code or "",
                    "location": p.location or "", "last_used": created.isoformat() if created else ""})
        if len(out) >= 25:
            break
    return out


# ---------------------------------------------------------------------------
# Order printouts (PDF / Word) — e.g. Business Development requests
# ---------------------------------------------------------------------------
def _order_access(db: Session, order: Order, user: User) -> bool:
    if order.requester_user_id == user.id or user.has_role("admin") or user.has_role("approver"):
        return True
    if db.query(OrderCollaborator).filter_by(order_id=order.id, user_id=user.id).first():
        return True
    return bool(order.project_id and access_service.project_membership(db, order.project_id, user.id))


def _load_order_for(db: Session, order_id: int, user: User) -> Order:
    order = db.query(Order).filter(Order.id == order_id, Order.deleted_at.is_(None)).first()
    if not order or not _order_access(db, order, user):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Order not found.")
    return order


def _print_out(row: OrderPrintFile) -> dict:
    return {"id": row.id, "filename": row.filename, "size_bytes": row.size_bytes or 0, "copies": row.copies,
            "sides": row.sides, "color_mode": row.color_mode, "paper": row.paper, "notes": row.notes or "",
            "printed": bool(row.printed), "printed_at": row.printed_at.isoformat() if row.printed_at else None,
            "created_at": row.created_at.isoformat() if row.created_at else None}


@router.post("/orders/{order_id}/print-files", status_code=201)
async def upload_print_file(order_id: int, file: UploadFile = File(...),
                            copies: int = Form(1), sides: str = Form("single"),
                            color_mode: str = Form("color"), paper: str = Form("Letter"),
                            notes: str = Form(""), db: Session = Depends(get_db),
                            user: User = Depends(get_current_user)):
    order = _load_order_for(db, order_id, user)
    privileged = user.has_role("admin") or user.has_role("approver")
    if order.status != "pending" and not privileged:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Printouts can only be added while the order is pending.")
    name = (file.filename or "printout.pdf").strip()
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if ext not in PRINT_EXTENSIONS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Upload a PDF or Word file (.pdf, .doc, .docx).")
    content = await file.read()
    if not content or len(content) > MAX_PRINT_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Each printout must be between 1 byte and 25 MB.")
    if sides not in {"single", "double"} or color_mode not in {"color", "bw"} or paper not in {"Letter", "Legal", "Tabloid"}:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown print setting.")
    row = OrderPrintFile(order_id=order.id, filename=name[:255], content_type=PRINT_EXTENSIONS[ext],
                         content=content, size_bytes=len(content), copies=max(1, min(int(copies or 1), 5000)),
                         sides=sides, color_mode=color_mode, paper=paper, notes=(notes or "")[:255])
    db.add(row)
    order.updated_at = utcnow()
    log_action(db, user_id=user.id, action="order.print_file_add", object_type="order", object_id=order.id,
               new_value={"filename": row.filename, "copies": row.copies, "sides": sides,
                          "color_mode": color_mode, "paper": paper}, source="browser")
    db.commit(); db.refresh(row)
    return _print_out(row)


@router.get("/orders/{order_id}/print-files/{file_id}", include_in_schema=False)
def download_print_file(order_id: int, file_id: int, db: Session = Depends(get_db),
                        user: User = Depends(get_current_user_flexible)):
    order = _load_order_for(db, order_id, user)
    row = db.query(OrderPrintFile).filter_by(id=file_id, order_id=order.id).first()
    if not row:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Printout not found.")
    safe = (row.filename or "printout.pdf").replace('"', "")
    return Response(content=row.content, media_type=row.content_type or "application/octet-stream",
                    headers={"Content-Disposition": f'attachment; filename="{safe}"',
                             "Cache-Control": "private, no-store"})


@router.put("/orders/{order_id}/print-files/{file_id}")
def update_print_file(order_id: int, file_id: int, body: PrintFileUpdate, db: Session = Depends(get_db),
                      user: User = Depends(get_current_user)):
    order = _load_order_for(db, order_id, user)
    row = db.query(OrderPrintFile).filter_by(id=file_id, order_id=order.id).first()
    if not row:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Printout not found.")
    privileged = user.has_role("admin") or user.has_role("approver")
    data = body.model_dump(exclude_none=True)
    if "printed" in data:
        if not privileged:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Only the warehouse can mark printouts as printed.")
        row.printed = bool(data.pop("printed"))
        row.printed_at = utcnow() if row.printed else None
    if data:
        if order.status != "pending" and not privileged:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Print settings can only change while the order is pending.")
        for k, v in data.items():
            setattr(row, k, v)
    order.updated_at = utcnow()
    db.commit(); db.refresh(row)
    return _print_out(row)


@router.delete("/orders/{order_id}/print-files/{file_id}", status_code=204)
def delete_print_file(order_id: int, file_id: int, db: Session = Depends(get_db),
                      user: User = Depends(get_current_user)):
    order = _load_order_for(db, order_id, user)
    row = db.query(OrderPrintFile).filter_by(id=file_id, order_id=order.id).first()
    if not row:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Printout not found.")
    privileged = user.has_role("admin") or user.has_role("approver")
    if order.status != "pending" and not privileged:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Printouts can only be removed while the order is pending.")
    log_action(db, user_id=user.id, action="order.print_file_remove", object_type="order", object_id=order.id,
               old_value={"filename": row.filename}, source="browser")
    db.delete(row)
    order.updated_at = utcnow()
    db.commit()
    return None


# ---------------------------------------------------------------------------
# Badge Studio storage — one document per event, plus a shared fallback
# ---------------------------------------------------------------------------
def _badge_scope(db: Session, scope: str, user: User) -> tuple[str, int | None]:
    scope = (scope or "shared").strip()
    if scope == "shared":
        return "shared", None
    if scope.startswith("project:") and scope[8:].isdigit():
        pid = int(scope[8:])
        project = db.query(Project).filter(Project.id == pid, Project.deleted_at.is_(None)).first()
        if not project or not access_service.can_view_project(db, user, project):
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found.")
        return f"project:{pid}", pid
    raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown badge project.")


@router.get("/badge-studio/scopes")
def badge_scopes(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    docs = {b.scope_key: b for b in db.query(BadgeProject).all()}
    projects = access_service.visible_projects_query(db, user).order_by(Project.event_date.desc(), Project.name).all()
    # Per-order private snapshots are inactive, so include events the user has ordered for too.
    own = (db.query(Project).join(Order, Order.project_id == Project.id)
           .filter(Order.requester_user_id == user.id, Order.deleted_at.is_(None), Project.deleted_at.is_(None))
           .order_by(Project.event_date.desc()).limit(50).all())
    seen, out = set(), []
    for p in list(projects) + list(own):
        if p.id in seen:
            continue
        seen.add(p.id)
        key = f"project:{p.id}"
        doc = docs.get(key)
        count = 0
        if doc and doc.document_json:
            try:
                count = len(json.loads(doc.document_json).get("attendees", []))
            except Exception:
                count = 0
        out.append({"scope": key, "name": p.name, "event_date": p.event_date or "",
                    "has_badges": bool(doc), "attendees": count,
                    "updated_at": doc.updated_at.isoformat() if doc and doc.updated_at else ""})
    shared = docs.get("shared")
    out.append({"scope": "shared", "name": "Shared badge project (no event)", "event_date": "",
                "has_badges": bool(shared), "attendees": 0,
                "updated_at": shared.updated_at.isoformat() if shared and shared.updated_at else ""})
    return out


@router.get("/badge-studio/document")
def get_badge_document(scope: str = Query("shared"), db: Session = Depends(get_db),
                       user: User = Depends(get_current_user)):
    key, _pid = _badge_scope(db, scope, user)
    row = db.query(BadgeProject).filter_by(scope_key=key).first()
    if not row or not row.document_json:
        return {"document": None, "revision": 0}
    return {"document": json.loads(row.document_json), "revision": row.revision}


@router.put("/badge-studio/document")
async def put_badge_document(payload: dict, scope: str = Query("shared"), db: Session = Depends(get_db),
                             user: User = Depends(get_current_user)):
    key, pid = _badge_scope(db, scope, user)
    document = payload.get("document")
    revision = payload.get("revision")
    if not isinstance(document, dict) or not isinstance(document.get("attendees"), list):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The badge project is invalid.")
    text = json.dumps(document, separators=(",", ":"))
    if len(text) > MAX_BADGE_DOC_CHARS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The badge project is too large.")
    row = db.query(BadgeProject).filter_by(scope_key=key).first()
    current = row.revision if row else 0
    if not isinstance(revision, int) or revision != current:
        raise HTTPException(status.HTTP_409_CONFLICT, "Badges changed in another window. Reload before saving.")
    if row is None:
        row = BadgeProject(scope_key=key, project_id=pid, document_json=text, revision=1,
                           updated_by_user_id=user.id)
        db.add(row)
    else:
        row.document_json = text
        row.revision = current + 1
        row.updated_by_user_id = user.id
    db.commit()
    return {"revision": row.revision}


# ---------------------------------------------------------------------------
# Reports (admin app "Reports" tab)
# ---------------------------------------------------------------------------
def _hours(a, b):
    a, b = _aware(a), _aware(b)
    if not a or not b:
        return None
    return max(0.0, (b - a).total_seconds() / 3600.0)


def _avg(values):
    vals = [v for v in values if v is not None]
    return round(sum(vals) / len(vals), 1) if vals else None


@router.get("/admin/reports/timing")
def report_timing(days: int = Query(90, ge=7, le=730), db: Session = Depends(get_db),
                  _u: User = Depends(require_role("admin", "approver"))):
    now = datetime.now(timezone.utc)
    since = now - timedelta(days=days)
    orders = (db.query(Order).options(joinedload(Order.approvals), joinedload(Order.requester),
                                      joinedload(Order.project))
              .filter(Order.deleted_at.is_(None)).all())
    rows = []
    for o in orders:
        created = _aware(o.created_at)
        approved_at = None
        for a in sorted(o.approvals, key=lambda x: x.created_at):
            if a.decision == "approved":
                approved_at = _aware(a.created_at)
                break
        rows.append({"order": o, "created": created, "approved": approved_at,
                     "pick": _aware(o.picking_started_at), "done": _aware(o.fulfilled_at)})
    in_window = [r for r in rows if r["created"] and r["created"] >= since]
    stage = {
        "approval_hours": _avg([_hours(r["created"], r["approved"]) for r in in_window]),
        "to_pick_hours": _avg([_hours(r["approved"], r["pick"]) for r in in_window]),
        "picking_hours": _avg([_hours(r["pick"], r["done"]) for r in in_window]),
        "total_hours": _avg([_hours(r["created"], r["done"]) for r in in_window]),
    }
    weeks = {}
    for r in in_window:
        wk = (r["created"] - timedelta(days=r["created"].weekday())).date().isoformat()
        w = weeks.setdefault(wk, {"week": wk, "orders": 0, "approval": [], "to_pick": [], "picking": [], "total": []})
        w["orders"] += 1
        w["approval"].append(_hours(r["created"], r["approved"]))
        w["to_pick"].append(_hours(r["approved"], r["pick"]))
        w["picking"].append(_hours(r["pick"], r["done"]))
        w["total"].append(_hours(r["created"], r["done"]))
    weekly = [{"week": w["week"], "orders": w["orders"], "approval_hours": _avg(w["approval"]),
               "to_pick_hours": _avg(w["to_pick"]), "picking_hours": _avg(w["picking"]),
               "total_hours": _avg(w["total"])} for w in sorted(weeks.values(), key=lambda x: x["week"])]

    def aging(status_name, since_key):
        out = []
        for r in rows:
            o = r["order"]
            if o.status != status_name:
                continue
            start = r[since_key] or r["created"]
            out.append({"id": o.id, "order_number": o.order_number or f"ORD-{o.id:06d}",
                        "requester": (o.requester.full_name or o.requester.username) if o.requester else "",
                        "project": o.project.name if o.project else "",
                        "ship_by": (o.project.ship_by_date if o.project else "") or "",
                        "waiting_hours": round(_hours(start, now) or 0, 1)})
        return sorted(out, key=lambda x: -x["waiting_hours"])[:25]

    counts = {}
    for r in in_window:
        counts[r["order"].status] = counts.get(r["order"].status, 0) + 1
    return {"days": days, "orders_in_window": len(in_window), "status_counts": counts,
            "averages": stage, "weekly": weekly,
            "aging": {"waiting_approval": aging("pending", "created"),
                      "waiting_pick": aging("approved", "approved"),
                      "being_picked": aging("picking", "pick")}}


def _demand_rows(db: Session, days: int) -> list[dict]:
    now = datetime.now(timezone.utc)
    since = now - timedelta(days=days)
    shipped = dict(db.query(OrderLine.item_id,
                            func.coalesce(func.sum(func.coalesce(OrderLine.qty_approved, OrderLine.qty_requested)), 0))
                   .join(Order, Order.id == OrderLine.order_id)
                   .filter(Order.status == "fulfilled", Order.deleted_at.is_(None), Order.fulfilled_at >= since)
                   .group_by(OrderLine.item_id).all())
    orders_count = dict(db.query(OrderLine.item_id, func.count(func.distinct(OrderLine.order_id)))
                        .join(Order, Order.id == OrderLine.order_id)
                        .filter(Order.status == "fulfilled", Order.deleted_at.is_(None), Order.fulfilled_at >= since)
                        .group_by(OrderLine.item_id).all())
    last_move = dict(db.query(InventoryTransaction.item_id, func.max(InventoryTransaction.created_at))
                     .group_by(InventoryTransaction.item_id).all())
    weeks = max(1.0, days / 7.0)
    out = []
    for item in db.query(Item).filter(Item.deleted_at.is_(None), Item.active.is_(True)).order_by(Item.code).all():
        if (item.review_status or "live") != "live":
            continue
        qty = int(item.qty_on_hand or 0)
        units = int(shipped.get(item.id, 0) or 0)
        per_week = units / weeks
        lm = _aware(last_move.get(item.id)) or _aware(item.created_at)
        idle_days = (now - lm).days if lm else None
        weeks_left = round(qty / per_week, 1) if per_week > 0 else None
        if qty > 0 and units == 0 and (idle_days is None or idle_days >= 180):
            flag = "dead_stock"
        elif weeks_left is not None and weeks_left < 4:
            flag = "reorder_soon"
        elif qty <= int(item.reorder_threshold or 0):
            flag = "low_stock"
        else:
            flag = ""
        out.append({"id": item.id, "code": item.code, "name": item.name, "brand": item.brand or "",
                    "category": item.category or "", "on_hand": qty, "shipped": units,
                    "orders": int(orders_count.get(item.id, 0) or 0), "per_week": round(per_week, 1),
                    "weeks_left": weeks_left, "last_movement": lm.date().isoformat() if lm else "",
                    "idle_days": idle_days, "reorder_threshold": int(item.reorder_threshold or 0),
                    "cost": float(item.cost or 0), "stock_value": round(qty * float(item.cost or 0), 2),
                    "flag": flag})
    return out


@router.get("/admin/reports/demand")
def report_demand(days: int = Query(90, ge=7, le=730), db: Session = Depends(get_db),
                  _u: User = Depends(require_role("admin", "approver"))):
    rows = _demand_rows(db, days)
    dead = [r for r in rows if r["flag"] == "dead_stock"]
    return {"days": days, "items": rows,
            "summary": {"items": len(rows), "units_shipped": sum(r["shipped"] for r in rows),
                        "dead_stock_items": len(dead),
                        "dead_stock_value": round(sum(r["stock_value"] for r in dead), 2),
                        "reorder_soon": sum(1 for r in rows if r["flag"] == "reorder_soon")}}


@router.get("/admin/reports/demand.xlsx", include_in_schema=False)
def report_demand_xlsx(days: int = Query(90, ge=7, le=730), db: Session = Depends(get_db),
                       _u: User = Depends(require_role("admin", "approver"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = _demand_rows(db, days)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Demand {days}d"
    headers = ["Part #", "Item", "Brand", "Category", "On hand", f"Shipped ({days}d)", "Orders",
               "Per week", "Weeks left", "Last movement", "Idle days", "Stock value", "Flag"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="A3277A")
    labels = {"dead_stock": "Dead stock", "reorder_soon": "Reorder soon", "low_stock": "Low stock", "": ""}
    for r in rows:
        ws.append([r["code"], r["name"], r["brand"], r["category"], r["on_hand"], r["shipped"], r["orders"],
                   r["per_week"], r["weeks_left"], r["last_movement"], r["idle_days"], r["stock_value"],
                   labels.get(r["flag"], r["flag"])])
    widths = [16, 34, 18, 16, 10, 14, 9, 10, 11, 14, 10, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="otiwarehouse-demand-{days}d.xlsx"'})
